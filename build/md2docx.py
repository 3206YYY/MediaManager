# -*- coding: utf-8 -*-
"""Markdown -> Word (.docx) 转换脚本。
支持：标题(#~####)、表格(GFM)、无序/有序列表、加粗/行内代码、分隔线、引用。
用法：python md2docx.py <input.md> <output.docx>
"""
import sys
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def set_cn_font(run, name="微软雅黑", size=None, bold=None, color=None):
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_runs_with_inline(paragraph, text, base_size, base_bold=False, base_color=None):
    """解析行内 Markdown（**加粗**、`代码`）并添加到段落。"""
    # 先处理 **加粗**
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            set_cn_font(run, size=base_size, bold=True, color=base_color)
        else:
            # 处理 `行内代码`
            for seg in re.split(r'(`[^`]+`)', part):
                if not seg:
                    continue
                if seg.startswith('`') and seg.endswith('`'):
                    run = paragraph.add_run(seg[1:-1])
                    set_cn_font(run, name="Consolas", size=base_size - 1, color=(0xC7, 0x25, 0x4E))
                else:
                    run = paragraph.add_run(seg)
                    set_cn_font(run, size=base_size, bold=base_bold, color=base_color)


def convert(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()

        # 空行
        if not line.strip():
            i += 1
            continue

        # 分隔线 ---
        if re.match(r'^\s*-{3,}\s*$', line):
            p = doc.add_paragraph()
            run = p.add_run('─' * 30)
            set_cn_font(run, color=(0xCC, 0xCC, 0xCC), size=8)
            i += 1
            continue

        # 标题
        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                p = doc.add_heading(level=0)
                run = p.add_run(text)
                set_cn_font(run, size=24, bold=True, color=(0x1F, 0x3B, 0x63))
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif level == 2:
                p = doc.add_heading(level=1)
                run = p.add_run(text)
                set_cn_font(run, size=17, bold=True, color=(0x1F, 0x3B, 0x63))
            elif level == 3:
                p = doc.add_heading(level=2)
                run = p.add_run(text)
                set_cn_font(run, size=14, bold=True, color=(0x2E, 0x5A, 0x88))
            else:
                p = doc.add_heading(level=3)
                run = p.add_run(text)
                set_cn_font(run, size=12, bold=True, color=(0x2E, 0x5A, 0x88))
            i += 1
            continue

        # 表格（GFM）
        if line.strip().startswith('|') and i + 1 < n and re.match(r'^\s*\|[\s:\-|]+\|\s*$', lines[i + 1]):
            # 收集表头
            header_cells = [c.strip() for c in line.strip().strip('|').split('|')]
            i += 2  # 跳过表头 + 分隔行
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header_cells))
            table.style = 'Light Grid Accent 1'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            # 表头
            for c, text in enumerate(header_cells):
                cell = table.cell(0, c)
                cell.text = ''
                p = cell.paragraphs[0]
                run = p.add_run(text)
                set_cn_font(run, size=10, bold=True, color=(0xFF, 0xFF, 0xFF))
            # 数据行
            for r, row in enumerate(rows):
                for c, text in enumerate(row):
                    cell = table.cell(r + 1, c)
                    cell.text = ''
                    p = cell.paragraphs[0]
                    add_runs_with_inline(p, text, base_size=10)
            doc.add_paragraph()
            continue

        # 无序列表
        m = re.match(r'^\s*[-*]\s+(.*)$', line)
        if m:
            text = m.group(1).strip()
            p = doc.add_paragraph(style='List Bullet')
            add_runs_with_inline(p, text, base_size=11)
            i += 1
            continue

        # 有序列表
        m = re.match(r'^\s*\d+\.\s+(.*)$', line)
        if m:
            text = m.group(1).strip()
            p = doc.add_paragraph(style='List Number')
            add_runs_with_inline(p, text, base_size=11)
            i += 1
            continue

        # 引用
        m = re.match(r'^\s*>\s*(.*)$', line)
        if m:
            text = m.group(1).strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs_with_inline(p, text, base_size=10, base_color=(0x88, 0x88, 0x88))
            i += 1
            continue

        # 普通段落
        add_runs_with_inline(doc.add_paragraph(), line.strip(), base_size=11)
        i += 1

    doc.save(docx_path)
    print(f"OK -> {docx_path}")


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print("用法：python md2docx.py <input.md> <output.docx>")
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
